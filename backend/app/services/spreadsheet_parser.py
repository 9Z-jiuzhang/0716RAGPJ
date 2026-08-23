"""Excel / CSV → 表格版面块（供多模态检索与 FAQ）。"""

from __future__ import annotations

import csv
import io
import logging
from typing import Any

from app.models.enums import ContentBlockType
from app.services.layout_blocks import ContentBlock, table_rows_to_html, table_rows_to_markdown
from app.services.parsers import _decode_bytes
from app.utils.exceptions import UnsupportedFileTypeError

logger = logging.getLogger(__name__)

# 单表过大时按行切分，避免超长 chunk；与版面表格切分策略一致
_MAX_ROWS_PER_BLOCK = 80
_MAX_SHEETS = 30
_MAX_COLS = 64


def extract_spreadsheet_blocks(content: bytes, file_type: str) -> list[ContentBlock]:
    """将 xlsx / xls / csv 解析为 Text(工作表名) + Table 块。"""
    ft = (file_type or "").lower().lstrip(".")
    if ft == "csv":
        return _extract_csv_blocks(content)
    if ft == "xlsx":
        return _extract_xlsx_blocks(content)
    if ft == "xls":
        return _extract_xls_blocks(content)
    raise UnsupportedFileTypeError(ft)


def spreadsheet_to_plain_text(content: bytes, file_type: str) -> str:
    """纯文本兜底：各表 Markdown 拼接。"""
    blocks = extract_spreadsheet_blocks(content, file_type)
    parts: list[str] = []
    for b in blocks:
        if b.block_type == ContentBlockType.TEXT.value:
            parts.append(b.content)
        elif b.block_type == ContentBlockType.TABLE.value:
            parts.append(b.markdown or b.content)
    return "\n\n".join(p for p in parts if (p or "").strip())


def _extract_csv_blocks(content: bytes) -> list[ContentBlock]:
    text = _decode_bytes(content)
    if not text.strip():
        return []
    # 去掉 BOM 残留
    if text.startswith("\ufeff"):
        text = text.lstrip("\ufeff")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    matrix = [_normalize_row(row) for row in reader]
    matrix = [r for r in matrix if any(c.strip() for c in r)]
    if not matrix:
        return []
    return _matrix_to_blocks(matrix, sheet_name="Sheet1")


def _extract_xlsx_blocks(content: bytes) -> list[ContentBlock]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise UnsupportedFileTypeError("xlsx(缺少 openpyxl 依赖)") from exc

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedFileTypeError(f"xlsx(无法打开: {exc})") from exc

    blocks: list[ContentBlock] = []
    try:
        for idx, name in enumerate(wb.sheetnames[:_MAX_SHEETS]):
            ws = wb[name]
            matrix: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_to_str(v) for v in (row or ())]
                if len(cells) > _MAX_COLS:
                    cells = cells[:_MAX_COLS]
                if any(c.strip() for c in cells):
                    matrix.append(cells)
            if not matrix:
                continue
            blocks.extend(_matrix_to_blocks(matrix, sheet_name=str(name) or f"Sheet{idx + 1}"))
    finally:
        wb.close()
    return blocks


def _extract_xls_blocks(content: bytes) -> list[ContentBlock]:
    """旧版 .xls：优先 xlrd；不可用时提示另存为 xlsx。"""
    try:
        import xlrd
    except ImportError as exc:
        raise UnsupportedFileTypeError("xls(请安装 xlrd，或另存为 .xlsx 后上传)") from exc
    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise UnsupportedFileTypeError(f"xls(无法打开: {exc})") from exc

    blocks: list[ContentBlock] = []
    for idx in range(min(book.nsheets, _MAX_SHEETS)):
        sheet = book.sheet_by_index(idx)
        matrix: list[list[str]] = []
        for r in range(sheet.nrows):
            cells = [_cell_to_str(sheet.cell_value(r, c)) for c in range(min(sheet.ncols, _MAX_COLS))]
            if any(c.strip() for c in cells):
                matrix.append(cells)
        if not matrix:
            continue
        blocks.extend(_matrix_to_blocks(matrix, sheet_name=sheet.name or f"Sheet{idx + 1}"))
    return blocks


def _matrix_to_blocks(matrix: list[list[str]], *, sheet_name: str) -> list[ContentBlock]:
    """首行作表头；过大时按行切分并重复表头。"""
    width = max(len(r) for r in matrix)
    norm = [r + [""] * (width - len(r)) for r in matrix]
    headers = [h.strip() or f"列{i + 1}" for i, h in enumerate(norm[0])]
    data_rows = norm[1:] if len(norm) > 1 else []
    # 只有一行：当作表头+空数据不如整表作为单列表值
    if not data_rows:
        data_rows = [headers]
        headers = [f"列{i + 1}" for i in range(len(headers))]

    out: list[ContentBlock] = []
    title = _text_block(f"工作表：{sheet_name}")
    if title:
        out.append(title)

    if not data_rows:
        return out

    for start in range(0, len(data_rows), _MAX_ROWS_PER_BLOCK):
        chunk_rows = data_rows[start : start + _MAX_ROWS_PER_BLOCK]
        tb = _table_block(headers, chunk_rows, sheet=sheet_name, row_offset=start)
        if tb:
            out.append(tb)
    return out


def _table_block(
    headers: list[str],
    rows: list[list[str]],
    **extra: Any,
) -> ContentBlock | None:
    md = table_rows_to_markdown(headers, rows)
    if not md.strip():
        return None
    merged = dict(extra)
    sheet = merged.get("sheet_name") or merged.get("sheet")
    merged.update(
        {
            "source_type": "spreadsheet",
            "sheet_name": sheet,
            "columns": headers,
            "rows": len(rows),
        }
    )
    return ContentBlock(
        block_type=ContentBlockType.TABLE.value,
        content=md,
        markdown=md,
        html=table_rows_to_html(headers, rows),
        extra=merged,
    )


def _text_block(text: str) -> ContentBlock | None:
    cleaned = (text or "").strip()
    if not cleaned:
        return None
    return ContentBlock(block_type=ContentBlockType.TEXT.value, content=cleaned)


def _normalize_row(row: list[Any]) -> list[str]:
    cells = [_cell_to_str(c) for c in (row or [])]
    if len(cells) > _MAX_COLS:
        return cells[:_MAX_COLS]
    return cells


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".") if "." in str(value) else str(value)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value).strip()
