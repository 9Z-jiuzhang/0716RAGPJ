"""多格式文本抽取。【对齐手册 §5.5 P0 格式】"""

from __future__ import annotations

import io
import logging
import re
import zlib
from pathlib import Path

from app.models.enums import DocumentFileType
from app.utils.exceptions import UnsupportedFileTypeError

logger = logging.getLogger(__name__)


def extract_text(filename: str, content: bytes, file_type: str) -> str:
    """从上传字节中抽取纯文本；失败时抛出明确错误，避免返回乱码。"""
    ft = file_type.lower().lstrip(".")
    try:
        if ft == DocumentFileType.TXT.value or ft == DocumentFileType.MD.value:
            return _decode_bytes(content)
        if ft == DocumentFileType.PDF.value:
            return _extract_pdf(content)
        if ft == DocumentFileType.DOCX.value:
            return _extract_docx(content)
        if ft == DocumentFileType.PPTX.value:
            return _extract_pptx(content)
        if ft == DocumentFileType.DOC.value:
            return _extract_doc(content)
    except UnsupportedFileTypeError:
        raise
    except Exception as exc:
        logger.exception("extract_text failed filename=%s type=%s", filename, ft)
        raise UnsupportedFileTypeError(f"{ft}(解析失败: {exc})") from exc
    raise UnsupportedFileTypeError(ft)


def _decode_bytes(content: bytes) -> str:
    """统一字节解码：优先 BOM/chardet，再回退常见中文与 Unicode 编码。"""
    if not content:
        return ""

    candidates: list[str] = []
    # BOM 优先，避免把 UTF-16/UTF-8-SIG 文件误判为乱码
    if content.startswith(b"\xff\xfe"):
        candidates.append("utf-16-le")
    elif content.startswith(b"\xfe\xff"):
        candidates.append("utf-16-be")
    elif content.startswith(b"\xef\xbb\xbf"):
        candidates.append("utf-8-sig")

    detected = _detect_encoding(content)
    if detected:
        candidates.append(detected)
    candidates.extend(
        [
            "utf-8",
            "utf-8-sig",
            "utf-16",
            "utf-16-le",
            "utf-16-be",
            "gb18030",
            "gbk",
            "gb2312",
            "big5",
            "latin-1",
        ]
    )

    seen: set[str] = set()
    best_fallback: str | None = None
    best_score = -1.0
    for enc in candidates:
        key = enc.lower().replace("_", "-")
        if key in seen:
            continue
        seen.add(key)
        try:
            text = content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
        score = _text_quality_score(text)
        if score > best_score:
            best_score = score
            best_fallback = text
        if not _looks_garbled(text):
            return text

    # 兜底：选质量最高的一次解码；中文正文足够时放行
    if best_fallback is not None and (
        best_score >= 0.45 or _cjk_ratio(best_fallback) >= 0.08
    ):
        return best_fallback

    text = content.decode("utf-8", errors="replace")
    if _looks_garbled(text) and _cjk_ratio(text) < 0.05:
        raise UnsupportedFileTypeError("txt(编码无法识别，请另存为 UTF-8 或 GBK)")
    return text


def _detect_encoding(content: bytes) -> str | None:
    sample = content[:65536]
    try:
        import chardet

        result = chardet.detect(sample) or {}
        enc = result.get("encoding")
        conf = float(result.get("confidence") or 0)
        if enc and conf >= 0.35:
            return str(enc)
    except Exception:
        logger.debug("chardet unavailable or failed", exc_info=True)
    return None


def _cjk_ratio(text: str) -> float:
    if not text:
        return 0.0
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    return cjk / max(len(text), 1)


def _text_quality_score(text: str) -> float:
    """越高越像可读正文（兼顾中英文）。"""
    if not text:
        return 0.0
    replacement = text.count("\ufffd")
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\t\r")
    printable_ratio = printable / max(len(text), 1)
    replacement_penalty = min(0.5, replacement / max(len(text), 1) * 4)
    return printable_ratio + _cjk_ratio(text) * 0.35 - replacement_penalty


def _looks_garbled(text: str) -> bool:
    if not text:
        return True
    replacement = text.count("\ufffd")
    if replacement > max(8, len(text) // 20):
        return True
    # 含较多汉字时放宽可打印比例，避免 UTF-16/混编码被误杀
    if _cjk_ratio(text) >= 0.08:
        printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\t\r")
        return printable / max(len(text), 1) < 0.45
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\t\r")
    return printable / max(len(text), 1) < 0.6


# PDF CID / 自定义编码乱码：/uni00000037 或 (cid:123)
_PDF_CID_TOKEN_RE = re.compile(r"(?:/uni[0-9a-fA-F]{4,8}|\(cid:\d+\))", re.IGNORECASE)

_RAG_PDF_VISION_PROMPT = (
    "你是 PDF 页面文字抽取助手。请用中文完整转写截图中的全部可读文字与图表信息，要求："
    "1) 保留标题、正文、图注、图例、坐标轴与关键数值；"
    "2) 图表数据尽量整理成 Markdown 表格；"
    "3) 提炼一句话结论；"
    "4) 只输出可检索的纯文本/Markdown，不要代码块围栏，不要道歉或解释过程。"
)


def is_unusable_pdf_text(text: str | None) -> bool:
    """判断 PDF 抽取结果是否为 CID/字形乱码或几乎不可检索。"""
    if not text or not str(text).strip():
        return True
    s = str(text)
    if _PDF_CID_TOKEN_RE.search(s):
        tokens = _PDF_CID_TOKEN_RE.findall(s)
        # 出现多个 uni/cid 标记，或标记占比过高 → 乱码
        if len(tokens) >= 3:
            return True
        stripped = _PDF_CID_TOKEN_RE.sub("", s)
        if len(stripped.strip()) < max(24, len(s) // 10):
            return True
    cjk = sum(1 for ch in s if "\u4e00" <= ch <= "\u9fff")
    if cjk >= 8:
        return False
    ascii_keep = sum(1 for ch in s if ch.isascii() and (ch.isalnum() or ch in " \n\t.%-+/"))
    non_ascii = sum(1 for ch in s if (not ch.isascii()) and ch not in "\n\t\r")
    # 纯英文数字文档可用；夹杂异常非 ASCII 且无中文则视为不可用
    if non_ascii == 0 and ascii_keep >= 20:
        return False
    return True


def _extract_pdf(content: bytes) -> str:
    """PDF 文本抽取：原生抽取 → 乱码则多模态看图转写。"""
    candidates: list[str] = []
    for name, extractor in (
        ("pypdf", _extract_pdf_pypdf),
        ("pymupdf", _extract_pdf_pymupdf),
    ):
        try:
            text = (extractor(content) or "").strip()
        except Exception as exc:
            logger.warning("pdf extract via %s failed: %s", name, exc)
            continue
        if not text:
            continue
        candidates.append(text)
        if not is_unusable_pdf_text(text):
            return text

    vision = ""
    try:
        vision = (_extract_pdf_via_vision(content) or "").strip()
    except Exception as exc:
        logger.warning("pdf vision fallback failed: %s", exc)
    if vision and not is_unusable_pdf_text(vision):
        return vision

    # 绝不把 /uni CID 乱码写入知识库污染检索
    if candidates and is_unusable_pdf_text(candidates[0]):
        raise UnsupportedFileTypeError(
            "pdf(页面文字无法抽取：字体缺少 ToUnicode/为 Type3 字形。"
            "已尝试多模态转写但仍失败，请检查 MARKITDOWN_LLM_* / LLM_API_KEY，"
            "或另存为可复制文本的 PDF 后重传)"
        )
    return vision or (candidates[0] if candidates else "")


def _extract_pdf_pypdf(content: bytes) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(content))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _extract_pdf_pymupdf(content: bytes) -> str:
    import fitz

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        parts = [(page.get_text("text") or "").strip() for page in doc]
        return "\n\n".join(p for p in parts if p)
    finally:
        doc.close()


def _extract_pdf_via_vision(content: bytes) -> str:
    """将每页栅格化后用多模态模型转写，供 CID 乱码 PDF 入库检索。"""
    from app.core.config import settings

    if not settings.MARKITDOWN_LLM_ENABLED:
        logger.warning("MARKITDOWN_LLM_ENABLED=false，跳过 PDF 看图转写")
        return ""

    # 延迟导入，避免与 markitdown_export 循环依赖
    from app.services.markitdown_export import (
        _llm_describe_png,
        _markitdown_llm_kwargs,
        _rasterize_pdf_pages,
    )

    kwargs = _markitdown_llm_kwargs()
    client = kwargs.get("llm_client")
    model = kwargs.get("llm_model")
    if client is None or not model:
        return ""

    pages = _rasterize_pdf_pages(content, zoom=2.0)
    if not pages:
        return ""

    sections: list[str] = []
    for idx, png in enumerate(pages, start=1):
        try:
            body = (_llm_describe_png(png, client=client, model=str(model), prompt=_RAG_PDF_VISION_PROMPT) or "").strip()
        except Exception as exc:
            logger.warning("pdf vision page %s failed: %s", idx, exc)
            body = ""
        if body:
            sections.append(f"## 第 {idx} 页\n\n{body}")
    return "\n\n".join(sections)


def _extract_docx(content: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs if p.text and p.text.strip())


def _extract_pptx(content: bytes) -> str:
    """抽取 PPTX 幻灯片文本（图表本身由 MarkItDown+LLM 在导出时描述）。"""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    parts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                    if cells:
                        texts.append(" | ".join(cells))
        if texts:
            parts.append(f"## 幻灯片 {idx}\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_doc_via_external(content: bytes) -> str | None:
    """优先用 antiword/catdoc 抽取（对中文 .doc 最稳）。"""
    import shutil
    import subprocess
    import tempfile

    tools = []
    if shutil.which("antiword"):
        tools.append(("antiword", ["antiword", "-m", "UTF-8.txt"]))
    if shutil.which("catdoc"):
        tools.append(("catdoc", ["catdoc", "-d", "utf-8"]))
    if not tools:
        return None

    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(content)
        path = tmp.name
    try:
        for name, cmd in tools:
            try:
                proc = subprocess.run(
                    [*cmd, path],
                    check=False,
                    capture_output=True,
                    timeout=30,
                )
                raw = proc.stdout or b""
                if not raw.strip():
                    continue
                text = raw.decode("utf-8", errors="ignore")
                cleaned = _clean_extracted_text(text)
                if cleaned and len(cleaned.strip()) >= 16 and not _looks_garbled(cleaned):
                    return cleaned
            except Exception:
                logger.warning("%s doc extract failed", name, exc_info=True)
    finally:
        try:
            Path(path).unlink(missing_ok=True)
        except Exception:
            pass
    return None


def _extract_doc(content: bytes) -> str:
    """旧版 Word .doc（OLE）抽取。

    python-docx 仅支持 .docx；此处按顺序尝试：
    1) 误标为 doc 的 docx
    2) antiword / catdoc（推荐，中文最稳）
    3) olefile 抽取 WordDocument 流 + 编码检测
    4) 原始字节流编码检测（严格校验，避免乱码冒充成功）
    """
    if content[:2] == b"PK":
        try:
            return _extract_docx(content)
        except Exception:
            logger.warning("doc marked as zip/docx but parse failed")

    external = _extract_doc_via_external(content)
    if external:
        return external

    ole_text = _extract_doc_via_ole(content)
    if ole_text and not _looks_garbled(ole_text) and len(ole_text.strip()) >= 16:
        # OLE 启发式可能夹杂噪声：若几乎没有汉字且原文像二进制，继续降级
        cjk = sum(1 for ch in ole_text if "\u4e00" <= ch <= "\u9fff")
        if cjk >= 8 or len(ole_text.strip()) >= 40:
            return ole_text

    # 非 OLE 或 OLE 失败：禁止把二进制当 latin-1 直接当正文
    try:
        text = _decode_bytes(content)
    except UnsupportedFileTypeError as exc:
        raise UnsupportedFileTypeError("doc(无法解析，请转换为 docx 后上传)") from exc
    cleaned = _clean_extracted_text(text)
    if len(cleaned.strip()) < 16 or _looks_garbled(cleaned):
        raise UnsupportedFileTypeError("doc(无法解析，请转换为 docx 后上传)")
    return cleaned


def _extract_doc_via_ole(content: bytes) -> str | None:
    try:
        import olefile
    except ImportError:
        logger.warning("olefile not installed; skip OLE .doc parse")
        return None

    if not olefile.isOleFile(io.BytesIO(content)):
        return None

    try:
        with olefile.OleFileIO(io.BytesIO(content)) as ole:
            stream_names = (
                "WordDocument",
                "1Table",
                "0Table",
                "RawText",
            )
            blobs: list[bytes] = []
            for name in stream_names:
                if ole.exists(name):
                    try:
                        blobs.append(ole.openstream(name).read())
                    except Exception:
                        continue
            if not blobs:
                return None
            # WordDocument 流中常夹杂 UTF-16LE 文本
            pieces: list[str] = []
            for blob in blobs:
                pieces.append(_extract_utf16_pieces(blob))
                pieces.append(_decode_bytes(blob) if _mostly_text_bytes(blob) else "")
            merged = _clean_extracted_text("\n".join(p for p in pieces if p))
            return merged or None
    except Exception:
        logger.warning("olefile doc extract failed", exc_info=True)
        return None


def _mostly_text_bytes(data: bytes) -> bool:
    if not data:
        return False
    sample = data[:4096]
    # 可打印 ASCII / 常见换行占比
    good = sum(1 for b in sample if 32 <= b <= 126 or b in (9, 10, 13))
    return good / max(len(sample), 1) > 0.75


def _is_useful_utf16_piece(piece: str) -> bool:
    """判断 OLE 抽出的 UTF-16 片段是否像正文（过滤字体/主题等噪声）。"""
    text = (piece or "").strip()
    if len(text) < 2:
        return False
    cjk = sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")
    if cjk >= 2:
        # 汉字占比过低多为二进制误读
        return cjk / max(len(text), 1) >= 0.25
    # 纯英文：允许较长句子，排除字体名
    lower = text.lower()
    if any(x in lower for x in ("times new roman", "arial", "cambria", "symbol", "theme", ".xml", "xmlns")):
        return False
    return len(text) >= 16 and " " in text and text.isascii()


def _extract_utf16_pieces(data: bytes) -> str:
    """从 WordDocument 二进制中提取可读的 UTF-16LE 片段（含中文）。"""
    # UTF-16LE：ASCII 可打印 或 CJK 统一汉字 U+4E00–U+9FFF
    pattern = re.compile(rb"(?:(?:[\x20-\x7e]\x00)|(?:[\x00-\xff][\x4e-\x9f])){4,}")
    chunks: list[str] = []

    def _collect(blob: bytes) -> None:
        for match in pattern.finditer(blob):
            try:
                piece = match.group().decode("utf-16-le", errors="ignore").strip()
            except Exception:
                continue
            if _is_useful_utf16_piece(piece):
                chunks.append(piece)

    _collect(data)
    # 再尝试 zlib 压缩块（部分 .doc 使用）
    for i in range(len(data) - 2):
        if data[i] == 0x78 and data[i + 1] in (0x01, 0x9C, 0xDA):
            try:
                inflated = zlib.decompress(data[i : i + 65536])
                _collect(inflated)
            except Exception:
                continue
    # 去重保序
    seen: set[str] = set()
    ordered: list[str] = []
    for c in chunks:
        if c in seen:
            continue
        seen.add(c)
        ordered.append(c)
    return _clean_extracted_text("\n".join(ordered))


def _clean_extracted_text(text: str) -> str:
    cleaned = "".join(ch if ch.isprintable() or ch in "\n\t\r" else " " for ch in text)
    lines = [re.sub(r"[ \t]{2,}", " ", line).strip() for line in cleaned.splitlines()]
    return "\n".join(line for line in lines if line)


def detect_file_type(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    return ext
